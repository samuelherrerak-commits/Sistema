"""Rellena el SISTEMA.xlsm con los asientos de datos.json, lo recalcula con LibreOffice y compara
la hoja Acumulados (saldo anterior / debe / haber / saldo por mes) contra lo que calculó el motor JS."""
import json, subprocess, sys, os, datetime, warnings, shutil
warnings.filterwarnings('ignore')
import openpyxl
src, datos, work = sys.argv[1], sys.argv[2], sys.argv[3]
os.makedirs(work, exist_ok=True)
d = json.load(open(datos))
wb = openpyxl.load_workbook(src)
mv, plan, ac = wb['Movimientos'], wb['Plan'], wb['Acumulados']
r = 3
for a in d['asientos']:
    for l in a['lineas']:
        mv.cell(r, 1, a['numero'])
        mv.cell(r, 3, datetime.datetime.strptime(a['fecha'], '%Y-%m-%d'))
        mv.cell(r, 4, int(l['cuenta']))
        mv.cell(r, 6, l['descripcion'])
        mv.cell(r, 7, l['debe']); mv.cell(r, 8, l['haber'])
        r += 1
print('filas escritas', r - 3)
codes = [str(plan.cell(i, 1).value) for i in range(4, 4 + 302)]
for i, c in enumerate(codes):
    if c in d['saldosIni']:
        ac.cell(4 + i, 6, d['saldosIni'][c])
p = os.path.join(work, 'in.xlsx'); wb.save(p)
subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', os.path.join(work, 'out'), p], check=True, capture_output=True)
wb2 = openpyxl.load_workbook(os.path.join(work, 'out', 'in.xlsx'), data_only=True)
a2 = wb2['Acumulados']
bad = 0; checked = 0
for i, c in enumerate(codes):
    row = 4 + i
    exp = d['acumulados'].get(c)
    for m in range(12):
        col = 25 + m * 4
        got = [a2.cell(row, col + k).value or 0 for k in range(4)]
        want = exp[m] if exp else [0, 0, 0, 0]
        checked += 1
        if any(abs(float(g) - w) > 0.005 for g, w in zip(got, want)):
            bad += 1
            if bad < 10: print('DIFERENCIA', c, 'mes', m + 1, 'excel', got, 'motor', want)
print(f'celdas cuenta-mes comparadas: {checked}, diferencias: {bad}')
sys.exit(1 if bad else 0)
